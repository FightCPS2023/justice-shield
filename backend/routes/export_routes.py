"""Export routes - PDF generation"""
import uuid
import io
import json
from datetime import datetime, timezone
from typing import Optional, List, Dict
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from database import get_db
from ai_mission import citation_log_for, log_to_markdown
from auth import get_current_user

router = APIRouter()

COURT_TEMPLATES: Dict[str, Dict[str, str]] = {
    "general_filing": {
        "name": "General Filing Cover",
        "body": "Court: [COURT NAME]\nCase No.: [CASE NUMBER]\nParties: [PARTY NAMES]\nTitle: [FILING TITLE]\n\nCOMES NOW the [Parent/Respondent], and respectfully submits this filing in support of their rights and the best interests of the child(ren)."
    },
    "motion_dismiss": {
        "name": "Motion to Dismiss",
        "body": "COMES NOW the [Parent/Respondent], and moves this Court to dismiss the petition due to constitutional violations, lack of jurisdiction, and insufficiency of evidence.\n\nGrounds:\n1. [Ground One]\n2. [Ground Two]\n3. [Ground Three]\n\nWHEREFORE, the parent respectfully requests dismissal and all just relief."
    },
    "motion_discovery": {
        "name": "Motion for Discovery",
        "body": "COMES NOW the [Parent/Respondent] and requests full discovery of all CPS records, communications, investigative notes, audio/video, and witness statements related to this case.\n\nRequested Categories:\n1. [Category One]\n2. [Category Two]\n3. [Category Three]"
    },
    "due_process": {
        "name": "Motion for Due Process Hearing",
        "body": "COMES NOW the [Parent/Respondent] and requests an immediate due process hearing regarding the removal and ongoing restrictions in this matter.\n\nBasis:\n- [Constitutional/Statutory Basis]\n- [Lack of Notice/Hearing]\n- [Emergency Exception Not Met]"
    },
    "emergency_motion": {
        "name": "Emergency Motion",
        "body": "COMES NOW the [Parent/Respondent] and moves for emergency relief due to imminent harm and ongoing constitutional violations.\n\nEmergency Facts:\n- [Fact One]\n- [Fact Two]\n- [Fact Three]"
    },
    "affidavit": {
        "name": "Affidavit / Declaration",
        "body": "I, [NAME], declare under penalty of perjury:\n\n1. [Statement One]\n2. [Statement Two]\n3. [Statement Three]\n\nExecuted on [DATE] in [CITY, STATE]."
    }
}

class ExportPayload(BaseModel):
    template_name: Optional[str] = "general_filing"
    template_text: Optional[str] = None
    editor_notes: Optional[str] = None
    include_analysis: Optional[bool] = True
    verified_only: Optional[bool] = False


async def _cite_guard_template(payload: "ExportPayload", case_id: str, user_id: str) -> Optional[str]:
    """Court export gate: run Cite-Guard on the filing text; verified-only mode strips non-verified authorities."""
    text = _get_template_content(payload.template_name, payload.template_text)
    if not text:
        return text
    log = await citation_log_for(get_db(), text, case_id, user_id)
    if payload.verified_only:
        return log["verified_only_text"] + log_to_markdown(log)
    if not log["release_grade"]:
        raise HTTPException(status_code=422, detail=f"{log['headline']} {log['counts'].get('UNVERIFIED', 0)} unverified and {log['counts'].get('STRIKE', 0)} struck authorities. Enable Verified-only mode or confirm the citations first.")
    return log["sanitized_text"] + log_to_markdown(log)


def _get_template_content(template_name: Optional[str], template_text: Optional[str]) -> str:
    if template_text:
        return template_text
    template = COURT_TEMPLATES.get(template_name or "", COURT_TEMPLATES["general_filing"])
    return template.get("body", "")


def _normalize_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        try:
            return json.dumps(value, indent=2)
        except Exception:
            return str(value)
    return str(value)


async def _fetch_case_export_data(db, case_id: str, user_id: str) -> dict:
    case = await db.cases.find_one({"_id": case_id, "user_id": user_id})
    if not case:
        raise HTTPException(status_code=404, detail="Case not found")

    violations = []
    async for v in db.violations.find({"case_id": case_id}):
        violations.append(v)

    timeline = []
    async for t in db.timeline_events.find({"case_id": case_id}).sort("event_date", 1):
        timeline.append(t)

    evidence = []
    async for e in db.evidence_items.find({"case_id": case_id}):
        evidence.append(e)

    documents = []
    async for d in db.documents.find({"case_id": case_id}):
        documents.append(d)

    notes = []
    async for n in db.notepad_entries.find({"case_id": case_id, "user_id": user_id}):
        notes.append(n)

    strategies = []
    async for s in db.strategies.find({"case_id": case_id}):
        strategies.append(s)

    communications = []
    async for c in db.communications.find({"case_id": case_id}):
        communications.append(c)

    court_filings = []
    async for f in db.court_filings.find({"case_id": case_id}):
        court_filings.append(f)

    return {
        "case": case,
        "violations": violations,
        "timeline": timeline,
        "evidence": evidence,
        "documents": documents,
        "notes": notes,
        "strategies": strategies,
        "communications": communications,
        "court_filings": court_filings
    }

@router.get("/api/export/notepad/all/pdf")
async def export_all_notes_to_pdf(
    case_id: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    """Export all notepad entries as single PDF"""
    db = get_db()
    from io import BytesIO
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.units import inch
    from fastapi.responses import StreamingResponse
    
    query = {"user_id": current_user["user_id"]}
    if case_id:
        query["case_id"] = case_id
    
    entries = []
    async for entry in db.notepad_entries.find(query).sort("created_at", -1):
        entries.append(entry)
    
    if not entries:
        raise HTTPException(status_code=404, detail="No notes found")
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.5*inch, bottomMargin=0.5*inch)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=20, spaceAfter=20)
    
    story = []
    story.append(Paragraph("THE CPS PUNISHER - Legal Research Notebook", title_style))
    story.append(Paragraph(f"Total Notes: {len(entries)}", styles['Normal']))
    story.append(Spacer(1, 20))
    
    for i, entry in enumerate(entries):
        if i > 0:
            story.append(PageBreak())
        
        story.append(Paragraph(entry.get("title", "Untitled"), styles['Heading2']))
        story.append(Paragraph(f"Category: {entry.get('category', 'general')} | Created: {entry.get('created_at', 'N/A')[:10]}", styles['Italic']))
        story.append(Spacer(1, 10))
        
        content = entry.get("content", "").replace("\n", "<br/>")
        content = content.replace("═", "-").replace("📋", "*").replace("📚", "*")
        content = content.replace("🏛️", "*").replace("📅", "*").replace("📝", "*")
        content = content.replace("🔗", "*").replace("⏱️", "*").replace("📜", "*")
        content = content.replace("⚖️", "*").replace("📖", "*").replace("📌", "*")
        
        story.append(Paragraph(content, styles['Normal']))
    
    doc.build(story)
    buffer.seek(0)
    
    filename = f"cps_punisher_all_notes.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/api/export/notepad/{entry_id}/pdf")
async def export_notepad_to_pdf(entry_id: str, current_user: dict = Depends(get_current_user)):
    """Export a notepad entry as PDF"""
    db = get_db()
    from io import BytesIO
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.units import inch
    from fastapi.responses import StreamingResponse
    
    entry = await db.notepad_entries.find_one({
        "_id": entry_id,
        "user_id": current_user["user_id"]
    })
    
    if not entry:
        raise HTTPException(status_code=404, detail="Note not found")
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.5*inch, bottomMargin=0.5*inch)
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=18, spaceAfter=20)
    body_style = ParagraphStyle('CustomBody', parent=styles['Normal'], fontSize=10, leading=14)
    
    story = []
    story.append(Paragraph(f"THE CPS PUNISHER - Legal Notepad", title_style))
    story.append(Paragraph(entry.get("title", "Untitled"), styles['Heading2']))
    story.append(Spacer(1, 12))
    
    # Convert content - handle special characters
    content = entry.get("content", "").replace("\n", "<br/>")
    content = content.replace("═", "-").replace("📋", "*").replace("📚", "*")
    content = content.replace("🏛️", "*").replace("📅", "*").replace("📝", "*")
    content = content.replace("🔗", "*").replace("⏱️", "*").replace("📜", "*")
    content = content.replace("⚖️", "*").replace("📖", "*").replace("📌", "*")
    
    story.append(Paragraph(content, body_style))
    story.append(Spacer(1, 20))
    story.append(Paragraph(f"Category: {entry.get('category', 'N/A')}", styles['Normal']))
    story.append(Paragraph(f"Created: {entry.get('created_at', 'N/A')}", styles['Normal']))
    
    doc.build(story)
    buffer.seek(0)
    
    filename = f"cps_punisher_note_{entry_id[:8]}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/api/export/case/{case_id}/pdf")
async def export_case_to_pdf(case_id: str, current_user: dict = Depends(get_current_user)):
    """Export full case summary as PDF"""
    db = get_db()
    from io import BytesIO
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from fastapi.responses import StreamingResponse
    
    # Get case
    case = await db.cases.find_one({"_id": case_id, "user_id": current_user["user_id"]})
    if not case:
        raise HTTPException(status_code=404, detail="Case not found")
    
    # Get related data
    violations = []
    async for v in db.violations.find({"case_id": case_id}):
        violations.append(v)
    
    timeline = []
    async for t in db.timeline_events.find({"case_id": case_id}).sort("event_date", 1):
        timeline.append(t)
    
    documents = []
    async for d in db.documents.find({"case_id": case_id}):
        documents.append(d)
    
    notes = []
    async for n in db.notepad_entries.find({"case_id": case_id, "user_id": current_user["user_id"]}):
        notes.append(n)
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.5*inch, bottomMargin=0.5*inch)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=20, spaceAfter=20)
    section_style = ParagraphStyle('Section', parent=styles['Heading2'], fontSize=14, spaceBefore=15, spaceAfter=10)
    
    story = []
    
    # Header
    story.append(Paragraph("THE CPS PUNISHER", title_style))
    story.append(Paragraph("CASE SUMMARY REPORT", styles['Heading2']))
    story.append(Spacer(1, 20))
    
    # Case Info
    story.append(Paragraph("CASE INFORMATION", section_style))
    case_info = [
        ["Case Title:", case.get("title", "N/A")],
        ["State:", case.get("state", "N/A")],
        ["County:", case.get("county", "N/A")],
        ["Case Number:", case.get("case_number", "N/A")],
        ["Status:", case.get("status", "active")],
        ["Created:", case.get("created_at", "N/A")[:10] if case.get("created_at") else "N/A"]
    ]
    t = Table(case_info, colWidths=[2*inch, 4*inch])
    t.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(t)
    story.append(Spacer(1, 15))
    
    # Violations
    if violations:
        story.append(Paragraph(f"VIOLATIONS IDENTIFIED ({len(violations)})", section_style))
        for v in violations:
            story.append(Paragraph(f"• {v.get('violation_type', 'Unknown')}: {v.get('description', '')[:100]}...", styles['Normal']))
        story.append(Spacer(1, 10))
    
    # Timeline
    if timeline:
        story.append(Paragraph(f"CASE TIMELINE ({len(timeline)} events)", section_style))
        for t in timeline[:10]:  # Limit to 10
            story.append(Paragraph(f"• {t.get('event_date', 'N/A')}: {t.get('title', 'Untitled')}", styles['Normal']))
        story.append(Spacer(1, 10))
    
    # Notes
    if notes:
        story.append(Paragraph(f"LEGAL NOTES ({len(notes)})", section_style))
        for n in notes[:5]:  # Limit to 5
            story.append(Paragraph(f"• {n.get('title', 'Untitled')} ({n.get('category', 'general')})", styles['Normal']))
    
    # Footer
    story.append(Spacer(1, 30))
    story.append(Paragraph("Generated by THE CPS PUNISHER - AI-Powered Legal Defense Platform", styles['Italic']))
    story.append(Paragraph("This document is for informational purposes only. Consult a licensed attorney.", styles['Italic']))
    
    doc.build(story)
    buffer.seek(0)
    
    filename = f"cps_case_{case_id[:8]}_summary.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.post("/api/export/case/{case_id}/documents/pdf")
async def export_case_documents_pdf(case_id: str, payload: ExportPayload, current_user: dict = Depends(get_current_user)):
    """Export all uploaded documents for a case as PDF."""
    db = get_db()
    data = await _fetch_case_export_data(db, case_id, current_user["user_id"])
    documents = data["documents"]
    case = data["case"]

    if not documents:
        raise HTTPException(status_code=404, detail="No documents found for this case")

    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.units import inch

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.5*inch, bottomMargin=0.5*inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=18, spaceAfter=16)
    section_style = ParagraphStyle('DocSection', parent=styles['Heading2'], fontSize=12, spaceBefore=12, spaceAfter=6)
    body_style = ParagraphStyle('DocBody', parent=styles['Normal'], fontSize=10, leading=14)

    template_content = await _cite_guard_template(payload, case_id, current_user["user_id"])
    editor_notes = _normalize_text(payload.editor_notes).strip()
    include_analysis = payload.include_analysis if payload.include_analysis is not None else True

    story = []
    story.append(Paragraph("THE CPS PUNISHER", title_style))
    story.append(Paragraph("DOCUMENT EXPORT", styles['Heading2']))
    story.append(Paragraph(f"Case: {case.get('title', 'Untitled')}", styles['Normal']))
    story.append(Spacer(1, 12))

    if template_content:
        story.append(Paragraph("COURT TEMPLATE", section_style))
        story.append(Paragraph(template_content.replace("\n", "<br/>") , body_style))
        story.append(Spacer(1, 10))

    if editor_notes:
        story.append(Paragraph("USER EDITS / CORRECTIONS", section_style))
        story.append(Paragraph(editor_notes.replace("\n", "<br/>") , body_style))
        story.append(Spacer(1, 10))

    for idx, doc_item in enumerate(documents, 1):
        if idx > 1:
            story.append(PageBreak())
        story.append(Paragraph(f"Document {idx}: {doc_item.get('title', doc_item.get('filename', 'Document'))}", section_style))
        story.append(Paragraph(f"Filename: {doc_item.get('filename', 'N/A')}", body_style))
        story.append(Paragraph(f"Uploaded: {doc_item.get('created_at', 'N/A')[:10] if doc_item.get('created_at') else 'N/A'}", body_style))
        story.append(Paragraph(f"Status: {(doc_item.get('analysis_status') or 'pending').upper()}", body_style))
        story.append(Paragraph(f"Type: {doc_item.get('file_type', 'N/A')}", body_style))
        story.append(Spacer(1, 6))

        if include_analysis and doc_item.get("analysis"):
            analysis = doc_item.get("analysis")
            summary = None
            if isinstance(analysis, dict):
                summary = analysis.get("summary") or analysis.get("overall_summary") or analysis.get("executive_summary")
            summary_text = _normalize_text(summary or analysis)
            story.append(Paragraph("Analysis Summary", section_style))
            story.append(Paragraph(summary_text[:2000].replace("\n", "<br/>") , body_style))

    doc.build(story)
    buffer.seek(0)

    filename = f"cps_case_{case_id[:8]}_documents.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.post("/api/export/case/{case_id}/documents/docx")
async def export_case_documents_docx(case_id: str, payload: ExportPayload, current_user: dict = Depends(get_current_user)):
    """Export all uploaded documents for a case as DOCX."""
    db = get_db()
    data = await _fetch_case_export_data(db, case_id, current_user["user_id"])
    documents = data["documents"]
    case = data["case"]

    if not documents:
        raise HTTPException(status_code=404, detail="No documents found for this case")

    from docx import Document

    template_content = await _cite_guard_template(payload, case_id, current_user["user_id"])
    editor_notes = _normalize_text(payload.editor_notes).strip()
    include_analysis = payload.include_analysis if payload.include_analysis is not None else True

    docx_file = Document()
    docx_file.add_heading("THE CPS PUNISHER - DOCUMENT EXPORT", level=1)
    docx_file.add_paragraph(f"Case: {case.get('title', 'Untitled')}")

    if template_content:
        docx_file.add_heading("Court Template", level=2)
        docx_file.add_paragraph(template_content)

    if editor_notes:
        docx_file.add_heading("User Edits / Corrections", level=2)
        docx_file.add_paragraph(editor_notes)

    for idx, doc_item in enumerate(documents, 1):
        docx_file.add_heading(f"Document {idx}: {doc_item.get('title', doc_item.get('filename', 'Document'))}", level=2)
        docx_file.add_paragraph(f"Filename: {doc_item.get('filename', 'N/A')}")
        docx_file.add_paragraph(f"Uploaded: {doc_item.get('created_at', 'N/A')[:10] if doc_item.get('created_at') else 'N/A'}")
        docx_file.add_paragraph(f"Status: {(doc_item.get('analysis_status') or 'pending').upper()}")
        docx_file.add_paragraph(f"Type: {doc_item.get('file_type', 'N/A')}")
        if include_analysis and doc_item.get("analysis"):
            analysis = doc_item.get("analysis")
            summary = None
            if isinstance(analysis, dict):
                summary = analysis.get("summary") or analysis.get("overall_summary") or analysis.get("executive_summary")
            summary_text = _normalize_text(summary or analysis)
            docx_file.add_heading("Analysis Summary", level=3)
            docx_file.add_paragraph(summary_text[:3000])

    buffer = io.BytesIO()
    docx_file.save(buffer)
    buffer.seek(0)

    filename = f"cps_case_{case_id[:8]}_documents.docx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.post("/api/export/case/{case_id}/preparation-kit/pdf")
async def export_case_preparation_kit_pdf(case_id: str, payload: ExportPayload, current_user: dict = Depends(get_current_user)):
    """Export case preparation kit (all data) as PDF."""
    db = get_db()
    data = await _fetch_case_export_data(db, case_id, current_user["user_id"])
    case = data["case"]

    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.units import inch

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.5*inch, bottomMargin=0.5*inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('KitTitle', parent=styles['Heading1'], fontSize=18, spaceAfter=16)
    section_style = ParagraphStyle('KitSection', parent=styles['Heading2'], fontSize=12, spaceBefore=12, spaceAfter=6)
    body_style = ParagraphStyle('KitBody', parent=styles['Normal'], fontSize=10, leading=14)

    template_content = await _cite_guard_template(payload, case_id, current_user["user_id"])
    editor_notes = _normalize_text(payload.editor_notes).strip()

    story = []
    story.append(Paragraph("THE CPS PUNISHER", title_style))
    story.append(Paragraph("CASE PREPARATION - ALL DATA", styles['Heading2']))
    story.append(Paragraph(f"Case: {case.get('title', 'Untitled')}", styles['Normal']))
    story.append(Spacer(1, 12))

    if template_content:
        story.append(Paragraph("COURT TEMPLATE", section_style))
        story.append(Paragraph(template_content.replace("\n", "<br/>") , body_style))
        story.append(Spacer(1, 10))

    if editor_notes:
        story.append(Paragraph("USER EDITS / CORRECTIONS", section_style))
        story.append(Paragraph(editor_notes.replace("\n", "<br/>") , body_style))
        story.append(Spacer(1, 10))

    story.append(Paragraph("CASE INFORMATION", section_style))
    story.append(Paragraph(f"State: {case.get('state', 'N/A')} | County: {case.get('county', 'N/A')} | Case No.: {case.get('case_number', 'N/A')}", body_style))
    story.append(Spacer(1, 6))

    sections = [
        ("VIOLATIONS", data["violations"], lambda v: f"{v.get('violation_type', 'Unknown')}: {v.get('description', '')[:200]}") ,
        ("TIMELINE", data["timeline"], lambda t: f"{t.get('event_date', 'N/A')[:10] if t.get('event_date') else 'N/A'} - {t.get('title', 'Untitled')}") ,
        ("EVIDENCE", data["evidence"], lambda e: f"{e.get('title', 'Evidence')} ({e.get('evidence_type', 'document')})") ,
        ("DOCUMENTS", data["documents"], lambda d: f"{d.get('filename', 'Document')} ({d.get('analysis_status', 'pending')})") ,
        ("LEGAL NOTES", data["notes"], lambda n: f"{n.get('title', 'Untitled')} - {n.get('category', 'general')}") ,
        ("STRATEGIES", data["strategies"], lambda s: f"{s.get('title', 'Strategy')} - {s.get('description', '')[:150]}") ,
        ("COMMUNICATIONS", data["communications"], lambda c: f"{c.get('contact_name', 'Contact')} - {c.get('summary', '')[:150]}") ,
        ("COURT FILINGS", data["court_filings"], lambda f: f"{f.get('title', 'Filing')} - {f.get('status', 'pending')}")
    ]

    for title, items, formatter in sections:
        story.append(PageBreak())
        story.append(Paragraph(title, section_style))
        if not items:
            story.append(Paragraph("No records available.", body_style))
            continue
        for item in items:
            story.append(Paragraph(f"• {formatter(item)}", body_style))

    doc.build(story)
    buffer.seek(0)

    filename = f"cps_case_{case_id[:8]}_all_data.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.post("/api/export/case/{case_id}/preparation-kit/docx")
async def export_case_preparation_kit_docx(case_id: str, payload: ExportPayload, current_user: dict = Depends(get_current_user)):
    """Export case preparation kit (all data) as DOCX."""
    db = get_db()
    data = await _fetch_case_export_data(db, case_id, current_user["user_id"])
    case = data["case"]

    from docx import Document

    template_content = await _cite_guard_template(payload, case_id, current_user["user_id"])
    editor_notes = _normalize_text(payload.editor_notes).strip()

    docx_file = Document()
    docx_file.add_heading("THE CPS PUNISHER - CASE PREPARATION (ALL DATA)", level=1)
    docx_file.add_paragraph(f"Case: {case.get('title', 'Untitled')}")
    docx_file.add_paragraph(f"State: {case.get('state', 'N/A')} | County: {case.get('county', 'N/A')} | Case No.: {case.get('case_number', 'N/A')}")

    if template_content:
        docx_file.add_heading("Court Template", level=2)
        docx_file.add_paragraph(template_content)

    if editor_notes:
        docx_file.add_heading("User Edits / Corrections", level=2)
        docx_file.add_paragraph(editor_notes)

    sections = [
        ("Violations", data["violations"], lambda v: f"{v.get('violation_type', 'Unknown')}: {v.get('description', '')[:200]}"),
        ("Timeline", data["timeline"], lambda t: f"{t.get('event_date', 'N/A')[:10] if t.get('event_date') else 'N/A'} - {t.get('title', 'Untitled')}"),
        ("Evidence", data["evidence"], lambda e: f"{e.get('title', 'Evidence')} ({e.get('evidence_type', 'document')})"),
        ("Documents", data["documents"], lambda d: f"{d.get('filename', 'Document')} ({d.get('analysis_status', 'pending')})"),
        ("Legal Notes", data["notes"], lambda n: f"{n.get('title', 'Untitled')} - {n.get('category', 'general')}"),
        ("Strategies", data["strategies"], lambda s: f"{s.get('title', 'Strategy')} - {s.get('description', '')[:150]}"),
        ("Communications", data["communications"], lambda c: f"{c.get('contact_name', 'Contact')} - {c.get('summary', '')[:150]}"),
        ("Court Filings", data["court_filings"], lambda f: f"{f.get('title', 'Filing')} - {f.get('status', 'pending')}")
    ]

    for title, items, formatter in sections:
        docx_file.add_heading(title, level=2)
        if not items:
            docx_file.add_paragraph("No records available.")
            continue
        for item in items:
            docx_file.add_paragraph(formatter(item), style='List Bullet')

    buffer = io.BytesIO()
    docx_file.save(buffer)
    buffer.seek(0)

    filename = f"cps_case_{case_id[:8]}_all_data.docx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ============== CASE PREPARATION KIT ==============

@router.get("/api/export/case/{case_id}/preparation-kit")
async def export_case_preparation_kit(case_id: str, current_user: dict = Depends(get_current_user)):
    """
    Generate comprehensive Court Preparation Kit PDF
    Includes: Case info, violations, timeline, evidence, documents, notes, strategies
    """
    db = get_db()
    from io import BytesIO
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, Image
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from fastapi.responses import StreamingResponse
    
    # Get case with ownership check
    case = await db.cases.find_one({"_id": case_id, "user_id": current_user["user_id"]})
    if not case:
        raise HTTPException(status_code=404, detail="Case not found")
    
    # Gather ALL case data
    violations = []
    async for v in db.violations.find({"case_id": case_id}).sort("date_occurred", -1):
        violations.append(v)
    
    timeline = []
    async for t in db.timeline_events.find({"case_id": case_id}).sort("event_date", 1):
        timeline.append(t)
    
    evidence = []
    async for e in db.evidence_items.find({"case_id": case_id}):
        evidence.append(e)
    
    documents = []
    async for d in db.documents.find({"case_id": case_id}):
        documents.append(d)
    
    notes = []
    async for n in db.notepad_entries.find({"case_id": case_id, "user_id": current_user["user_id"]}):
        notes.append(n)
    
    strategies = []
    async for s in db.strategies.find({"case_id": case_id}):
        strategies.append(s)
    
    communications = []
    async for c in db.communications.find({"case_id": case_id}).sort("date", -1):
        communications.append(c)
    
    court_filings = []
    async for f in db.court_filings.find({"case_id": case_id}).sort("filed_date", -1):
        court_filings.append(f)
    
    # Create PDF
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, 
        pagesize=letter, 
        topMargin=0.75*inch, 
        bottomMargin=0.75*inch,
        leftMargin=0.75*inch,
        rightMargin=0.75*inch
    )
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'KitTitle', 
        parent=styles['Heading1'], 
        fontSize=24, 
        spaceAfter=10,
        alignment=TA_CENTER,
        textColor=colors.darkred
    )
    subtitle_style = ParagraphStyle(
        'KitSubtitle',
        parent=styles['Normal'],
        fontSize=12,
        alignment=TA_CENTER,
        textColor=colors.gray,
        spaceAfter=30
    )
    section_style = ParagraphStyle(
        'Section', 
        parent=styles['Heading2'], 
        fontSize=16, 
        spaceBefore=20, 
        spaceAfter=12,
        textColor=colors.darkblue,
        borderWidth=1,
        borderColor=colors.darkblue,
        borderPadding=5
    )
    subsection_style = ParagraphStyle(
        'Subsection',
        parent=styles['Heading3'],
        fontSize=12,
        spaceBefore=10,
        spaceAfter=6,
        textColor=colors.black
    )
    body_style = ParagraphStyle(
        'KitBody',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        spaceAfter=6
    )
    bullet_style = ParagraphStyle(
        'KitBullet',
        parent=styles['Normal'],
        fontSize=10,
        leading=12,
        leftIndent=20,
        spaceAfter=4
    )
    
    story = []
    
    # ===== COVER PAGE =====
    story.append(Spacer(1, 1*inch))
    story.append(Paragraph("THE CPS PUNISHER", title_style))
    story.append(Paragraph("COURT PREPARATION KIT", ParagraphStyle('SubTitle', parent=title_style, fontSize=18, textColor=colors.black)))
    story.append(Spacer(1, 0.5*inch))
    story.append(Paragraph(f"Case: {case.get('title', 'Untitled Case')}", subtitle_style))
    story.append(Paragraph(f"State: {case.get('state', 'N/A')} | County: {case.get('county', 'N/A')}", subtitle_style))
    story.append(Paragraph(f"Case Number: {case.get('case_number', 'N/A')}", subtitle_style))
    story.append(Spacer(1, 0.5*inch))
    
    # Stats summary
    stats_data = [
        ["DOCUMENT SUMMARY", ""],
        ["Total Violations Identified:", str(len(violations))],
        ["Timeline Events:", str(len(timeline))],
        ["Evidence Items:", str(len(evidence))],
        ["Documents Analyzed:", str(len(documents))],
        ["Legal Notes:", str(len(notes))],
        ["Defense Strategies:", str(len(strategies))],
        ["Communications Logged:", str(len(communications))],
        ["Court Filings:", str(len(court_filings))],
    ]
    stats_table = Table(stats_data, colWidths=[3*inch, 2*inch])
    stats_table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
        ('FONTNAME', (0, 1), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 1), (-1, -1), 10),
        ('ALIGN', (1, 0), (1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(stats_table)
    story.append(Spacer(1, 0.5*inch))
    
    generation_date = datetime.now(timezone.utc).strftime("%B %d, %Y at %H:%M UTC")
    story.append(Paragraph(f"Generated: {generation_date}", ParagraphStyle('Date', parent=styles['Italic'], alignment=TA_CENTER)))
    story.append(PageBreak())
    
    # ===== TABLE OF CONTENTS =====
    story.append(Paragraph("TABLE OF CONTENTS", section_style))
    toc_items = [
        "1. Case Information",
        "2. Identified Violations",
        "3. Case Timeline",
        "4. Evidence Inventory",
        "5. Document Analysis",
        "6. Legal Research Notes",
        "7. Defense Strategies",
        "8. Communication Log",
        "9. Court Filings",
        "10. Legal Disclaimers"
    ]
    for item in toc_items:
        story.append(Paragraph(item, body_style))
    story.append(PageBreak())
    
    # ===== 1. CASE INFORMATION =====
    story.append(Paragraph("1. CASE INFORMATION", section_style))
    
    case_info_data = [
        ["Field", "Value"],
        ["Case Title", case.get("title", "N/A")],
        ["Case Number", case.get("case_number", "N/A")],
        ["State", case.get("state", "N/A")],
        ["County", case.get("county", "N/A")],
        ["Status", case.get("status", "active").upper()],
        ["Date Opened", case.get("date_opened", "N/A")],
        ["Caseworker", case.get("caseworker_name", "N/A")],
        ["Judge", case.get("judge_name", "N/A")],
        ["Court", case.get("court_name", "N/A")],
        ["Attorney", case.get("attorney_name", "N/A")],
    ]
    
    if case.get("children_names"):
        case_info_data.append(["Children", ", ".join(case.get("children_names", []))])
    
    case_table = Table(case_info_data, colWidths=[2*inch, 4.5*inch])
    case_table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BACKGROUND', (0, 0), (-1, 0), colors.darkblue),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 1), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    story.append(case_table)
    
    if case.get("description"):
        story.append(Spacer(1, 15))
        story.append(Paragraph("Case Description:", subsection_style))
        story.append(Paragraph(case.get("description", ""), body_style))
    
    story.append(PageBreak())
    
    # ===== 2. VIOLATIONS =====
    story.append(Paragraph("2. IDENTIFIED VIOLATIONS", section_style))
    
    if violations:
        story.append(Paragraph(f"Total violations identified: {len(violations)}", body_style))
        story.append(Spacer(1, 10))
        
        for i, v in enumerate(violations, 1):
            story.append(Paragraph(f"Violation #{i}: {v.get('violation_type', 'Unknown')}", subsection_style))
            
            v_data = [
                ["Severity:", v.get("severity", "medium").upper()],
                ["Date Occurred:", v.get("date_occurred", "N/A")],
                ["Description:", v.get("description", "No description")[:500]],
            ]
            if v.get("evidence_notes"):
                v_data.append(["Evidence Notes:", v.get("evidence_notes", "")[:300]])
            
            for row in v_data:
                story.append(Paragraph(f"<b>{row[0]}</b> {row[1]}", bullet_style))
            story.append(Spacer(1, 10))
    else:
        story.append(Paragraph("No violations have been documented yet.", body_style))
    
    story.append(PageBreak())
    
    # ===== 3. TIMELINE =====
    story.append(Paragraph("3. CASE TIMELINE", section_style))
    
    if timeline:
        story.append(Paragraph(f"Total events: {len(timeline)}", body_style))
        story.append(Spacer(1, 10))
        
        timeline_data = [["Date", "Event", "Type"]]
        for t in timeline:
            timeline_data.append([
                t.get("event_date", "N/A")[:10] if t.get("event_date") else "N/A",
                t.get("title", "Untitled")[:50],
                t.get("event_type", "general")
            ])
        
        timeline_table = Table(timeline_data, colWidths=[1.2*inch, 3.8*inch, 1.5*inch])
        timeline_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('BACKGROUND', (0, 0), (-1, 0), colors.darkblue),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))
        story.append(timeline_table)
    else:
        story.append(Paragraph("No timeline events have been documented yet.", body_style))
    
    story.append(PageBreak())
    
    # ===== 4. EVIDENCE =====
    story.append(Paragraph("4. EVIDENCE INVENTORY", section_style))
    
    if evidence:
        story.append(Paragraph(f"Total evidence items: {len(evidence)}", body_style))
        story.append(Spacer(1, 10))
        
        for i, e in enumerate(evidence, 1):
            story.append(Paragraph(f"Evidence #{i}: {e.get('title', 'Untitled')}", subsection_style))
            story.append(Paragraph(f"<b>Type:</b> {e.get('evidence_type', 'document')}", bullet_style))
            story.append(Paragraph(f"<b>Source:</b> {e.get('source', 'N/A')}", bullet_style))
            story.append(Paragraph(f"<b>Date Collected:</b> {e.get('date_collected', 'N/A')}", bullet_style))
            story.append(Paragraph(f"<b>Description:</b> {e.get('description', '')[:300]}", bullet_style))
            story.append(Spacer(1, 8))
    else:
        story.append(Paragraph("No evidence items have been catalogued yet.", body_style))
    
    story.append(PageBreak())
    
    # ===== 5. DOCUMENTS =====
    story.append(Paragraph("5. DOCUMENT ANALYSIS", section_style))
    
    if documents:
        story.append(Paragraph(f"Total documents analyzed: {len(documents)}", body_style))
        story.append(Spacer(1, 10))
        
        for i, d in enumerate(documents, 1):
            story.append(Paragraph(f"Document #{i}: {d.get('filename', d.get('title', 'Untitled'))}", subsection_style))
            story.append(Paragraph(f"<b>Type:</b> {d.get('document_type', 'general')}", bullet_style))
            story.append(Paragraph(f"<b>Uploaded:</b> {d.get('created_at', 'N/A')[:10] if d.get('created_at') else 'N/A'}", bullet_style))
            
            if d.get("ai_analysis"):
                analysis = d.get("ai_analysis", "")[:500]
                story.append(Paragraph(f"<b>AI Analysis:</b> {analysis}...", bullet_style))
            story.append(Spacer(1, 8))
    else:
        story.append(Paragraph("No documents have been uploaded for analysis yet.", body_style))
    
    story.append(PageBreak())
    
    # ===== 6. LEGAL NOTES =====
    story.append(Paragraph("6. LEGAL RESEARCH NOTES", section_style))
    
    if notes:
        # Group by category
        categories = {}
        for n in notes:
            cat = n.get("category", "general")
            if cat not in categories:
                categories[cat] = []
            categories[cat].append(n)
        
        for cat, cat_notes in categories.items():
            story.append(Paragraph(f"{cat.upper()} ({len(cat_notes)} notes)", subsection_style))
            for n in cat_notes[:5]:  # Limit to 5 per category
                story.append(Paragraph(f"• <b>{n.get('title', 'Untitled')}</b>", bullet_style))
                if n.get("source_type"):
                    story.append(Paragraph(f"  Source: {n.get('source_type')}", ParagraphStyle('Source', parent=bullet_style, fontSize=9, textColor=colors.gray)))
            story.append(Spacer(1, 8))
    else:
        story.append(Paragraph("No legal research notes have been saved yet.", body_style))
    
    story.append(PageBreak())
    
    # ===== 7. STRATEGIES =====
    story.append(Paragraph("7. DEFENSE STRATEGIES", section_style))
    
    if strategies:
        for i, s in enumerate(strategies, 1):
            story.append(Paragraph(f"Strategy #{i}: {s.get('title', 'Untitled')}", subsection_style))
            story.append(Paragraph(f"<b>Type:</b> {s.get('strategy_type', 'general')}", bullet_style))
            content = s.get("content", "")[:800]
            # Clean content for PDF
            content = content.replace("\n", " ").replace("•", "-")
            story.append(Paragraph(content, body_style))
            story.append(Spacer(1, 10))
    else:
        story.append(Paragraph("No defense strategies have been generated yet.", body_style))
    
    story.append(PageBreak())
    
    # ===== 8. COMMUNICATIONS =====
    story.append(Paragraph("8. COMMUNICATION LOG", section_style))
    
    if communications:
        story.append(Paragraph(f"Total communications logged: {len(communications)}", body_style))
        story.append(Spacer(1, 10))
        
        comm_data = [["Date", "Type", "Contact", "Summary"]]
        for c in communications[:20]:  # Limit to 20
            comm_data.append([
                c.get("date", "N/A")[:10] if c.get("date") else "N/A",
                c.get("communication_type", "N/A")[:15],
                c.get("contact_name", "N/A")[:20],
                c.get("summary", "")[:40] + "..."
            ])
        
        comm_table = Table(comm_data, colWidths=[1*inch, 1.2*inch, 1.5*inch, 2.8*inch])
        comm_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('BACKGROUND', (0, 0), (-1, 0), colors.darkblue),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
        ]))
        story.append(comm_table)
    else:
        story.append(Paragraph("No communications have been logged yet.", body_style))
    
    story.append(PageBreak())
    
    # ===== 9. COURT FILINGS =====
    story.append(Paragraph("9. COURT FILINGS", section_style))
    
    if court_filings:
        for i, f in enumerate(court_filings, 1):
            story.append(Paragraph(f"Filing #{i}: {f.get('title', 'Untitled')}", subsection_style))
            story.append(Paragraph(f"<b>Type:</b> {f.get('filing_type', 'N/A')}", bullet_style))
            story.append(Paragraph(f"<b>Filed Date:</b> {f.get('filed_date', 'N/A')}", bullet_style))
            story.append(Paragraph(f"<b>Status:</b> {f.get('status', 'filed')}", bullet_style))
            story.append(Spacer(1, 8))
    else:
        story.append(Paragraph("No court filings have been recorded yet.", body_style))
    
    story.append(PageBreak())
    
    # ===== 10. DISCLAIMERS =====
    story.append(Paragraph("10. LEGAL DISCLAIMERS", section_style))
    
    disclaimers = [
        "This document is generated by THE CPS PUNISHER, an AI-powered legal research and case management tool.",
        "This document is for INFORMATIONAL PURPOSES ONLY and does not constitute legal advice.",
        "Always consult with a licensed attorney before taking any legal action.",
        "The accuracy of information in this document depends on the data entered by the user.",
        "AI-generated analyses and strategies should be reviewed by qualified legal counsel.",
        "This document may contain sensitive personal information and should be handled confidentially.",
        "THE CPS PUNISHER and its creators are not responsible for any actions taken based on this document."
    ]
    
    for d in disclaimers:
        story.append(Paragraph(f"• {d}", bullet_style))
    
    story.append(Spacer(1, 30))
    story.append(Paragraph("— END OF COURT PREPARATION KIT —", ParagraphStyle('End', parent=styles['Italic'], alignment=TA_CENTER)))
    
    # Build PDF
    doc.build(story)
    buffer.seek(0)
    
    # Generate filename
    case_title_clean = "".join(c for c in case.get("title", "case")[:20] if c.isalnum() or c in " -_").strip().replace(" ", "_")
    filename = f"CPS_Punisher_Court_Kit_{case_title_clean}_{case_id[:8]}.pdf"
    
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )